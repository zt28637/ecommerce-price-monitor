"""
============================================================
③ 竞品数据反馈 — 竞品日报导出
============================================================
作用：读 sku_cleaned Excel → 生成带图片的竞品日报 Excel

功能：
  1. 同一商品的多行 SKU 自动合并商品级列
  2. 嵌入商品主图（D 列）和 SKU 属性图（G 列）
  3. 表头加粗 + 浅蓝底色 + 列宽适配
  4. 价格列居中，文字列垂直居中

输入：product_clean.py 输出的 13 列清洗 Excel（runtime/product_clean/）
输出：runtime/reports/ 下的日报 Excel

调用方式（影刀）：
  report.main("清洗Excel路径", "报表输出路径")
============================================================
"""

import json
import os
from io import BytesIO

import openpyxl
import requests
from PIL import Image as PILImage
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# 列布局配置（1-based = Excel 列号）
# ============================================================

# 合并列 — 同一商品的这些列会合并成一个单元格
# A=商品ID  B=商品名称  C=店铺名称  D=商品主图
# K=销量    L=采集时间  M=采集方式
MERGE_COLS = (1, 2, 3, 4, 11, 12, 13)

# 图片嵌入的列
COL_MAIN = 4   # D 列 — 商品主图（每个商品只嵌一次）
COL_SKU  = 7   # G 列 — SKU 图片（每行都嵌）

# 图片缩放到这个尺寸以内
IMG_MAX_W, IMG_MAX_H = 88, 76

# ============================================================
# 表头与格式
# ============================================================

# 输入列名 → 报表显示名（缩短中文名）
HEADER_RENAME = {
    "商品主图链接": "商品主图",
    "SKU主图链接":  "SKU图片",
    "SKUID":       "SKUID",
}

# 报表 13 列表头
REPORT_HEADERS = [
    "商品ID",   "商品名称",   "店铺名称",   "商品主图",
    "SKUID",    "SKU名称",   "SKU图片",
    "实际售价",  "原价",       "活动价格",
    "销量",      "采集时间",   "采集方式",
]

# 每列的宽度（Excel 字符单位）
_COL_WIDTHS = {
    1: 14, 2: 28, 3: 16, 4: 36, 5: 14,
    6: 24, 7: 36, 8: 12, 9: 10, 10: 10,
    11: 10, 12: 18, 13: 10,
}

# 价格 + 销量列：水平居中
_PRICE_COLS = {8, 9, 10, 11}


# ============================================================
# 工具函数
# ============================================================

def _abspath(path):
    """规范化路径：绝对路径 + 去外层引号"""
    return os.path.normpath(os.path.abspath(str(path).strip().strip("\"'")))


def _download_image(url):
    """
    下载网络图片 → 统一转 PNG（openpyxl 不支持 webp，阿里 CDN 不管 URL 后缀实际返 webp）
    失败（超时/404/非图片）返回 None，不阻断流程
    """
    url = str(url or "").strip()
    if not url.startswith("http"):
        return None
    try:
        resp = requests.get(url, timeout=15,
                           headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()

        # 统一用 PIL 转 PNG：无论原始格式是什么，openpyxl 只认 PNG/JPG/GIF
        pil_img = PILImage.open(BytesIO(resp.content))
        if pil_img.mode in ("RGBA", "P", "LA"):
            pil_img = pil_img.convert("RGBA")
        else:
            pil_img = pil_img.convert("RGB")
        png_data = BytesIO()
        pil_img.save(png_data, format="PNG")
        return png_data
    except Exception:
        return None


def _add_image(worksheet, row, col, url=None):
    """
    在指定单元格嵌入图片
    成功 → 清空单元格文字（只显示图片），返回 True
    失败 → 保留原文字，返回 False
    """
    url = str(url if url is not None else worksheet.cell(row, col).value or "").strip()
    if not url.startswith("http"):
        return False

    image_data = _download_image(url)
    if not image_data:
        return False

    img = XLImage(image_data)
    # 按比例缩放，不放大小图
    scale = min(IMG_MAX_W / img.width, IMG_MAX_H / img.height, 1.0)
    img.width = max(1, int(img.width * scale))
    img.height = max(1, int(img.height * scale))

    cell_ref = f"{get_column_letter(col)}{row}"
    worksheet.add_image(img, cell_ref)
    worksheet.cell(row, col).value = None
    return True


# ============================================================
# 合并 + 嵌图（核心逻辑）
# ============================================================

def _merge_and_embed(worksheet, start_row, last_row):
    """
    双指针遍历数据行：嵌图 + 合并商品级列

    算法说明：
      以"商品ID 变化"为块边界逐行扫描
      - 块首行 → 嵌入 D 列商品主图（每商品只嵌一次）
      - 每行   → 嵌入 G 列 SKU 图
      - 块结束 → 合并 A/B/C/D/K/L/M 列的单元格

    返回: (主图成功数, SKU图成功数)
    """
    main_ok = 0
    sku_ok = 0
    block_start = start_row

    for r in range(start_row, last_row + 1):
        # ---- 块首行嵌主图 ----
        if r == block_start:
            if _add_image(worksheet, r, COL_MAIN):
                main_ok += 1

        # ---- 每行嵌 SKU 图 ----
        if _add_image(worksheet, r, COL_SKU):
            sku_ok += 1

        # ---- 判断块边界 ----
        is_last = (r == last_row)
        current_id = worksheet.cell(r, 1).value
        next_id = worksheet.cell(r + 1, 1).value if not is_last else None

        if is_last or next_id != current_id:
            # 块结束，合并该块内的商品级列
            if r > block_start:
                for col in MERGE_COLS:
                    worksheet.merge_cells(
                        start_row=block_start, start_column=col,
                        end_row=r, end_column=col,
                    )
            block_start = r + 1

    return main_ok, sku_ok


# ============================================================
# 格式美化
# ============================================================

def _format_sheet(worksheet, last_row):
    """设置列宽、行高、对齐方式、表头样式"""
    # 表头
    worksheet.row_dimensions[1].height = 24
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    for col, width in _COL_WIDTHS.items():
        worksheet.column_dimensions[get_column_letter(col)].width = width
        cell = worksheet.cell(1, col)
        cell.font = Font(bold=True, size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for r in range(2, last_row + 1):
        worksheet.row_dimensions[r].height = 80
        for c in range(1, 14):
            cell = worksheet.cell(r, c)
            if c in _PRICE_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


# ============================================================
# 主流程
# ============================================================

def export_report(excel_path, output_path=None, keyword="", timestamp=""):
    """
    读清洗 Excel → 生成带图片的竞品日报

    参数:
        excel_path  : product_clean.py 输出的 13 列 Excel（runtime/product_clean/）
        output_path : 日报输出路径。不传则自动生成 runtime/reports/{关键词}_{时间}.xlsx
        keyword     : 搜索关键词（用于默认文件名）
        timestamp   : 采集时间（用于默认文件名）

    返回:
        {"ok": True, "filepath": "...", "rows": N, "products": N,
         "images_main": N, "images_sku": N}
    """
    excel_path = _abspath(excel_path)
    if not os.path.isfile(excel_path):
        return {"ok": False, "error": f"找不到清洗 Excel: {excel_path}"}

    # ---- 默认输出路径 ----
    if not output_path:
        if not keyword:
            basename = os.path.splitext(os.path.basename(excel_path))[0]
            # 从输入文件名提取关键词（SKU_小天才Z6pro_xxx → 小天才Z6pro）
            clean_name = basename
            if basename.startswith("SKU_"):
                clean_name = basename[4:]
            keyword = clean_name.split("_")[0] if "_" in clean_name else clean_name
        if not timestamp:
            from datetime import datetime
            timestamp = datetime.now().strftime("%Y年%m月%d日%H时%M分")
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        output_path = os.path.join(project_root, "runtime", "reports",
                                   f"{keyword}_{timestamp}.xlsx")

    output_path = _abspath(output_path)

    # ---- 只读打开清洗 Excel ----
    wb_in = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws_in = wb_in.active
    all_rows = list(ws_in.iter_rows(values_only=True))
    wb_in.close()

    # 分离表头和数据
    if all_rows:
        # 表头映射："商品主图链接" → "商品主图"
        headers = [HEADER_RENAME.get(str(h), str(h)) for h in all_rows[0]]
        data_rows = all_rows[1:]
    else:
        headers = list(REPORT_HEADERS)
        data_rows = []

    # ---- 写报表 ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竞品动态"
    ws.append(headers)

    for row in data_rows:
        if row and row[0] is not None:
            ws.append(list(row))

    last_row = ws.max_row
    main_imgs = sku_imgs = 0
    if last_row > 1:
        _format_sheet(ws, last_row)
        main_imgs, sku_imgs = _merge_and_embed(ws, 2, last_row)

    # ---- 保存 ----
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    wb.close()

    # ---- 统计 ----
    product_ids = {row[0] for row in data_rows if row and row[0] is not None}

    return {
        "ok": True,
        "filepath": output_path,
        "rows": len(data_rows),
        "products": len(product_ids),
        "images_main": main_imgs,
        "images_sku": sku_imgs,
    }


def _safe_path(p):
    """修复影刀传参时被 Python 转义的路径（如 \1 → \x01，还原为 \1）"""
    if p is None:
        return None
    s = str(p)
    for i in range(1, 10):
        s = s.replace(chr(i), '\\' + str(i))
    s = s.replace('\r', '\\r')
    s = s.replace('\n', '\\n')
    s = s.replace('\f', '\\f')
    s = s.replace('\v', '\\v')
    return s


def main(excel_path, output_path=None, keyword="", timestamp=""):
    """
    影刀入口函数

    用法：
        excel_path  = product_clean Excel 路径（runtime/product_clean/）
        output_path = 日报输出路径（可选）
        keyword     = glv['搜索关键词']
        timestamp   = glv['当前时间']
    """
    excel_path = _safe_path(excel_path)
    output_path = _safe_path(output_path)
    return export_report(excel_path, output_path=output_path,
                         keyword=keyword, timestamp=timestamp)


# ============================================================
# 本地调试
# ============================================================
if __name__ == "__main__":
    EXCEL = (
        r"F:\Programs\RPA自动化\1-淘宝竞品价格监控自动化"
        r"\project\runtime\product_clean\小天才Z6pro_2026年07月20日15时00分.xlsx"
    )
    OUTPUT = (
        r"F:\Programs\RPA自动化\1-淘宝竞品价格监控自动化"
        r"\project\runtime\reports\竞品日报_测试.xlsx"
    )
    print(json.dumps(export_report(EXCEL, OUTPUT), ensure_ascii=False, indent=2))
